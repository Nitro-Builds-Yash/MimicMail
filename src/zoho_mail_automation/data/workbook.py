"""Read, validate, and persist recipient rows in CSV/XLSX files."""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

EMAIL_RE = re.compile(r"^[^\s@]+@[^\s@]+\.[^\s@]+$")


@dataclass(slots=True)
class Recipient:
    row_number: int
    name: str
    email: str
    title: str


class RecipientWorkbook:
    def __init__(self, path: Path, columns: dict[str, str]):
        self.path = path
        self.columns = columns
        self._rows: list[dict[str, Any]] = []
        self._headers: list[str] = []
        self._sheet = None
        self._workbook = None
        self._load()

    def _load(self) -> None:
        if self.path.suffix.lower() == ".csv":
            with self.path.open("r", encoding="utf-8-sig", newline="") as stream:
                reader = csv.DictReader(stream)
                self._headers = list(reader.fieldnames or [])
                self._rows = [dict(row) for row in reader]
        else:
            self._workbook = load_workbook(self.path)
            self._sheet = self._workbook.active
            self._headers = [str(cell.value).strip() if cell.value is not None else "" for cell in self._sheet[1]]
            for row in self._sheet.iter_rows(min_row=2, values_only=True):
                self._rows.append({header: value for header, value in zip(self._headers, row) if header})

        required = [self.columns[k] for k in ("name", "email", "title")]
        missing = [name for name in required if name not in self._headers]
        if missing:
            raise ValueError(f"Workbook is missing required column(s): {', '.join(missing)}")
        for optional in ("status", "sent_at"):
            name = self.columns[optional]
            if name not in self._headers:
                self._headers.append(name)
                for row in self._rows:
                    row[name] = ""
                if self._sheet is not None:
                    self._sheet.cell(row=1, column=len(self._headers), value=name)

    def valid_recipients(self) -> tuple[list[Recipient], list[tuple[int, str]]]:
        recipients: list[Recipient] = []
        skipped: list[tuple[int, str]] = []
        for index, row in enumerate(self._rows, start=2):
            status = str(row.get(self.columns["status"]) or "").strip().upper()
            if status == "SENT":
                skipped.append((index, "already sent"))
                continue
            address = str(row.get(self.columns["email"]) or "").strip()
            if not EMAIL_RE.fullmatch(address):
                skipped.append((index, "invalid or missing email"))
                continue
            name = str(row.get(self.columns["name"]) or "").strip()
            title = str(row.get(self.columns["title"]) or "").strip()
            if not name or not title:
                skipped.append((index, "missing name or title"))
                continue
            recipients.append(Recipient(index, name, address, title))
        return recipients, skipped

    def mark_sent(self, row_number: int) -> None:
        row = self._rows[row_number - 2]
        stamp = datetime.now(timezone.utc).astimezone().isoformat(timespec="seconds")
        row[self.columns["status"]] = "SENT"
        row[self.columns["sent_at"]] = stamp
        if self._sheet is not None:
            status_col = self._headers.index(self.columns["status"]) + 1
            sent_at_col = self._headers.index(self.columns["sent_at"]) + 1
            self._sheet.cell(row=row_number, column=status_col, value="SENT")
            self._sheet.cell(row=row_number, column=sent_at_col, value=stamp)
            self._workbook.save(self.path)
        else:
            with self.path.open("w", encoding="utf-8", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=self._headers)
                writer.writeheader()
                writer.writerows(self._rows)

    def close(self) -> None:
        if self._workbook is not None:
            self._workbook.close()
